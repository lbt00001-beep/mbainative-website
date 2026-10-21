"""Lectura de libros y documentos, troceado y búsqueda de fragmentos relevantes."""
from __future__ import annotations

import hashlib
import io
import logging
import re
import threading
import unicodedata
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from sqlalchemy import update

from . import ai, config
from .db import Chunk, Document, OcrCache, SessionLocal, Subject, Topic

log = logging.getLogger("pupitre.ingest")

IMAGE_TYPES = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp",
               ".heic": "image/heic"}
ACCEPTED = {".pdf", ".docx", ".txt", ".md", ".csv", *IMAGE_TYPES}


def ext_of(name: str) -> str:
    m = re.search(r"\.[A-Za-z0-9]+$", name or "")
    return m.group(0).lower() if m else ""


def _set(doc_id: int, **fields) -> None:
    with SessionLocal() as db:
        d = db.get(Document, doc_id)
        if d:
            for k, v in fields.items():
                setattr(d, k, v)
            db.commit()


# ---------------------------------------------------------------------------
# Análisis previo (sin IA): qué páginas hay que leer con IA y cuánto costaría
# ---------------------------------------------------------------------------

_cancel: set[int] = set()  # documentos cuya lectura se ha pedido detener


def file_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def stored_path(doc_id: int, filename: str) -> Path:
    config.UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    return config.UPLOAD_DIR / f"{doc_id}{ext_of(filename)}"


def cached_pages(fhash: str) -> set[int]:
    with SessionLocal() as db:
        return {p for (p,) in db.query(OcrCache.page).filter(OcrCache.file_hash == fhash)}


def analyze(filename: str, data: bytes) -> dict:
    """Cuenta páginas y detecta las escaneadas (sin texto), que son las que se leen con IA."""
    ext = ext_of(filename)
    if ext == ".pdf":
        import pymupdf
        with pymupdf.open(stream=data, filetype="pdf") as pdf:
            n = pdf.page_count
            scanned = [i + 1 for i, page in enumerate(pdf) if len((page.get_text("text") or "").strip()) < 80]
        return {"pages": n, "scanned": scanned}
    if ext in IMAGE_TYPES:
        return {"pages": 1, "scanned": [1]}
    return {"pages": 1, "scanned": []}


def parse_pages(text: str) -> list[int]:
    return [int(x) for x in (text or "").split(",") if x.strip().isdigit()]


def estimate(scanned: list[int], cached: set[int], page_from: int, page_to: int) -> dict:
    todo = [p for p in scanned if page_from <= p <= page_to and p not in cached]
    return {"to_read": len(todo), "cost": round(len(todo) * config.OCR_COST_PER_PAGE, 2),
            "minutes": max(1, round(len(todo) * 2.5 / max(1, config.OCR_WORKERS) / 60))}


# ---------------------------------------------------------------------------
# Extracción de texto
# ---------------------------------------------------------------------------

def _normalize_image(data: bytes, mime: str) -> tuple[bytes, str]:
    """Reduce fotos grandes del móvil para abaratar la lectura (máx. 2000 px de lado)."""
    try:
        import pymupdf
        pix = pymupdf.Pixmap(data)
        if max(pix.width, pix.height) > 2000:
            f = 2000 / max(pix.width, pix.height)
            pix = pymupdf.Pixmap(pix, int(pix.width * f), int(pix.height * f), None)
        if pix.alpha:
            pix = pymupdf.Pixmap(pix, 0)
        return pix.tobytes("jpeg", jpg_quality=85), "image/jpeg"
    except Exception:  # noqa: BLE001 - si no se puede convertir, se envía tal cual
        return data, mime


def _ocr_pages(doc_id: int, fhash: str, todo: list[tuple[int, callable]], family_id: int) -> tuple[dict, bool]:
    """Lee con IA las páginas pendientes, guardando cada una en la caché al momento.
    Devuelve ({página: texto}, detenido)."""
    out: dict[int, str] = {}
    stopped = [False]

    def work(item):
        num, render = item
        if doc_id in _cancel:
            stopped[0] = True
            return num, None
        try:
            img, mime = render()
            text = ai.ocr_image(img, mime, family_id=family_id, doc_id=doc_id)
        except ai.AiError as e:
            if "límite diario" in str(e) or "saldo" in str(e) or "clave" in str(e).lower():
                _cancel.add(doc_id)  # sin crédito o sin clave no tiene sentido seguir
                stopped[0] = True
                _set(doc_id, error=str(e))
            log.warning("OCR página %s: %s", num, e)
            return num, None
        except Exception:  # noqa: BLE001 - una página defectuosa no tumba el documento entero
            log.exception("Página %s: no se pudo preparar o leer", num)
            return num, None
        with SessionLocal() as db:
            if not db.query(OcrCache).filter_by(file_hash=fhash, page=num).first():
                db.add(OcrCache(file_hash=fhash, page=num, text=text))
            db.execute(update(Document).where(Document.id == doc_id).values(ocr_done=Document.ocr_done + 1))
            db.commit()
        return num, text

    with ThreadPoolExecutor(max_workers=config.OCR_WORKERS) as ex:
        for num, text in ex.map(work, todo):
            if text is not None:
                out[num] = text
    return out, stopped[0] or doc_id in _cancel


def extract_pages(doc: Document, data: bytes, family_id: int) -> tuple[list[tuple[int, str]], bool]:
    """[(página, texto)] del rango elegido y si se detuvo antes de terminar."""
    ext = ext_of(doc.filename)
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return [(1, "\n".join(parts))], False
    if ext in (".txt", ".md", ".csv"):
        for enc in ("utf-8", "latin-1"):
            try:
                return [(1, data.decode(enc))], False
            except UnicodeDecodeError:
                continue
        raise ValueError("No se puede leer el texto")
    cache = cached_pages(doc.file_hash)
    with SessionLocal() as db:
        cached_text = {c.page: c.text for c in db.query(OcrCache).filter(OcrCache.file_hash == doc.file_hash)}
    if ext in IMAGE_TYPES:
        if 1 in cache:
            return [(1, cached_text[1])], False
        img = _normalize_image(data, IMAGE_TYPES[ext])
        _set(doc.id, ocr_total=1, ocr_done=0)
        got, stopped = _ocr_pages(doc.id, doc.file_hash, [(1, lambda: img)], family_id)
        return ([(1, got[1])] if 1 in got else []), stopped
    if ext != ".pdf":
        raise ValueError("Formato no admitido")

    import pymupdf
    pdf = pymupdf.open(stream=data, filetype="pdf")
    first = max(1, doc.page_from or 1)
    last = min(pdf.page_count, doc.page_to or pdf.page_count)
    scanned = set(parse_pages(doc.scanned_pages))
    pages: dict[int, str] = {}
    todo = []
    for num in range(first, last + 1):
        if num not in scanned:
            pages[num] = pdf[num - 1].get_text("text") or ""
        elif num in cache:
            pages[num] = cached_text[num]
        else:
            todo.append(num)
    todo = todo[: config.MAX_OCR_PAGES]
    in_range_scanned = [p for p in scanned if first <= p <= last]
    _set(doc.id, ocr_total=len(in_range_scanned), ocr_done=len(in_range_scanned) - len(todo))

    pdf_lock = threading.Lock()  # PyMuPDF no admite varios hilos sobre el mismo documento

    def renderer(num):
        def render():
            with pdf_lock:
                pix = pdf[num - 1].get_pixmap(dpi=120)
                return pix.tobytes("jpeg", jpg_quality=75), "image/jpeg"
        return render

    stopped = False
    if todo:
        got, stopped = _ocr_pages(doc.id, doc.file_hash, [(n, renderer(n)) for n in todo], family_id)
        pages.update({k: v for k, v in got.items() if v.strip()})
    pdf.close()
    return sorted((k, v) for k, v in pages.items() if v.strip()), stopped


# ---------------------------------------------------------------------------
# Troceado
# ---------------------------------------------------------------------------

def make_chunks(pages: list[tuple[int, str]], size: int = 1400, overlap: int = 200) -> list[tuple[int, str]]:
    out = []
    for page, text in pages:
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        if not text:
            continue
        start = 0
        while start < len(text):
            end = min(len(text), start + size)
            if end < len(text):
                cut = text.rfind("\n", start + size // 2, end)
                if cut == -1:
                    cut = text.rfind(". ", start + size // 2, end)
                end = cut + 1 if cut != -1 else end
            piece = text[start:end].strip()
            if len(piece) > 40:
                out.append((page, piece))
            if end >= len(text):
                break
            start = max(end - overlap, start + 1)
    return out


def segments(pages: list[tuple[int, str]], size: int) -> list[str]:
    out, cur, n = [], [], 0
    for p, t in pages:
        piece = f"[pág. {p}] {t}"
        if cur and n + len(piece) > size:
            out.append("\n\n".join(cur))
            cur, n = [], 0
        cur.append(piece[:size])
        n += len(piece)
    if cur:
        out.append("\n\n".join(cur))
    return out


def pdf_outline(data: bytes) -> str:
    """Índice interno (marcadores) del PDF, si lo tiene."""
    try:
        import pymupdf
        with pymupdf.open(stream=data, filetype="pdf") as pdf:
            toc = pdf.get_toc(simple=True)
        return "\n".join(f"{'  ' * (lvl - 1)}{title} (pág. {page})" for lvl, title, page in toc[:150] if lvl <= 3)
    except Exception:  # noqa: BLE001
        return ""


# ---------------------------------------------------------------------------
# Proceso completo de un documento (se ejecuta en segundo plano)
# ---------------------------------------------------------------------------

def start(doc_id: int, family_id: int) -> None:
    _cancel.discard(doc_id)
    _set(doc_id, status="procesando", progress=2, error="", note="")
    threading.Thread(target=process_document, args=(doc_id, family_id), daemon=True).start()


def request_stop(doc_id: int) -> None:
    _cancel.add(doc_id)


def process_document(doc_id: int, family_id: int) -> None:
    try:
        with SessionLocal() as db:
            doc = db.get(Document, doc_id)
            if doc is None:
                return
            db.expunge(doc)
        path = Path(doc.stored_file) if doc.stored_file else None
        if not path or not path.exists():
            _set(doc_id, status="detenido", note="El archivo ya no está en el servidor. Vuelve a subirlo: "
                                                  "las páginas ya leídas no se vuelven a pagar.")
            return
        data = path.read_bytes()
        pages, stopped = extract_pages(doc, data, family_id)
        if stopped:
            with SessionLocal() as db:
                d = db.get(Document, doc_id)
                if d:
                    d.status = "detenido"
                    d.note = (f"Detenido: {d.ocr_done} de {d.ocr_total} páginas escaneadas leídas. "
                              "Lo ya leído no se vuelve a pagar.")
                    db.commit()
            return
        chunks = make_chunks(pages)
        if not chunks:
            _set(doc_id, status="error", error="No se ha encontrado texto legible en el archivo.")
            return
        with SessionLocal() as db:
            d = db.get(Document, doc_id)
            if d is None:
                return
            for i, (page, text) in enumerate(chunks):
                db.add(Chunk(document_id=doc_id, subject_id=d.subject_id, idx=i, page=page, text=text))
            d.chars = sum(len(t) for _, t in chunks)
            d.progress = 90
            subject = db.get(Subject, d.subject_id)
            existing = [t.title for t in subject.topics]
            course = subject.student.course
            subject_name = subject.name
            subject_id = subject.id
            db.commit()

        toc = pdf_outline(data) if ext_of(doc.filename) == ".pdf" else ""
        topics = []
        # Libros completos: se analizan por tramos (unos 90.000 caracteres, ~60 páginas) para que
        # ningún tema se quede fuera. Cada tramo conoce los temas ya encontrados para no repetirlos.
        segs = segments(pages, 90_000)
        for k, seg in enumerate(segs):
            if len(existing) + len(topics) >= 40:
                break
            hint = (f"ÍNDICE DEL LIBRO (úsalo como guía de los temas):\n{toc}\n\n" if toc else "")
            found = ai.extract_topics(subject_name, course, hint + seg,
                                      existing + [t["title"] for t in topics], family_id=family_id, doc_id=doc_id)
            topics.extend(found)
            _set(doc_id, progress=min(99, 90 + int(9 * (k + 1) / max(1, len(segs)))))
        with SessionLocal() as db:
            pos = db.query(Topic).filter_by(subject_id=subject_id).count()
            for t in topics:
                db.add(Topic(subject_id=subject_id, title=t["title"], summary=t["summary"],
                             keywords=t["keywords"], position=pos))
                pos += 1
            d = db.get(Document, doc_id)
            d.status, d.progress, d.stored_file = "listo", 100, ""
            read = f"{d.ocr_total} páginas escaneadas leídas con IA · " if d.ocr_total else ""
            rng = f"páginas {d.page_from}-{d.page_to} · " if d.page_from and d.page_to else ""
            d.note = f"{rng}{read}{len(topics)} temas nuevos · coste {d.ai_cost:.2f} $".replace(".", ",")
            db.commit()
        path.unlink(missing_ok=True)  # el archivo original no se conserva
    except ai.AiError as e:
        _set(doc_id, status="error", error=str(e))
    except Exception as e:  # noqa: BLE001
        log.exception("Error procesando documento %s", doc_id)
        _set(doc_id, status="error", error=f"No se pudo leer el archivo ({type(e).__name__}).")
    finally:
        _cancel.discard(doc_id)


# ---------------------------------------------------------------------------
# Búsqueda de fragmentos relevantes (BM25 ligero, sin embeddings)
# ---------------------------------------------------------------------------

STOP = set("""a al algo como con de del donde el ella ellas ellos en entre era es esta este esto estos fue ha han hay
la las le les lo los mas me mi muy no nos o para pero por que se si sin sobre son su sus te tiene tu un una uno unos
y ya yo cual cuales cuando tambien porque cada otro otra otros otras puede pueden ser hacer tema""".split())


def tokens(text: str) -> list[str]:
    t = unicodedata.normalize("NFKD", text.lower())
    t = "".join(c for c in t if not unicodedata.combining(c))
    return [w[:7] for w in re.findall(r"[a-z0-9]{3,}", t) if w not in STOP]


def retrieve(subject_id: int, query: str, max_chars: int = 7000) -> str:
    with SessionLocal() as db:
        rows = db.query(Chunk.text, Chunk.page).filter(Chunk.subject_id == subject_id).all()
    if not rows:
        return ""
    q = Counter(tokens(query))
    if not q:
        return "\n\n".join(r.text for r in rows[:4])[:max_chars]
    docs = [Counter(tokens(r.text)) for r in rows]
    n = len(docs)
    df = Counter()
    for d in docs:
        df.update(set(d))
    avg = sum(sum(d.values()) for d in docs) / n or 1
    import math
    scored = []
    for i, d in enumerate(docs):
        L = sum(d.values()) or 1
        s = 0.0
        for w, qf in q.items():
            if w in d:
                idf = math.log(1 + (n - df[w] + 0.5) / (df[w] + 0.5))
                tf = d[w]
                s += idf * (tf * 2.2) / (tf + 1.2 * (0.25 + 0.75 * L / avg)) * qf
        scored.append((s, i))
    scored.sort(reverse=True)
    out, total = [], 0
    for s, i in scored:
        if s <= 0 and out:
            break
        t = rows[i].text
        if total + len(t) > max_chars:
            break
        out.append(f"[pág. {rows[i].page}] {t}")
        total += len(t)
    return "\n\n".join(out)


def topic_context(topic: Topic, max_chars: int = 7000) -> str:
    return retrieve(topic.subject_id, f"{topic.title} {topic.title} {topic.keywords} {topic.summary}", max_chars)
